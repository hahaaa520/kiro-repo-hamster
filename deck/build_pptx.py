# -*- coding: utf-8 -*-
"""代码仓鼠 · 由 slide-0x.png 合成 16:9 PPTX（含演讲备注 & 淡入切换）"""
import os
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.oxml.ns import qn
from lxml import etree

BASE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(BASE, "shots")
OUT = os.path.join(BASE, "代码仓鼠-路演.pptx")

SLIDES = [
    ("slide-01.png", """【开场 · 约 30 秒】
大家好，我们是「代码仓鼠」。先记住一句话：你写代码，它吃饭。
它不是工具，是一只住在你 GitHub 仓库里的电子宠物：
输入用户名，3 秒领养；你提交，它吃饭；你摆烂，它会饿、会 emo，
14 天它会离家出走。今天先给大家看一只饿了七天的仓鼠。"""),
    ("slide-02.png", """【共鸣 · 约 40 秒】
每个开发者的备忘录里，都躺着几个“有空就做”的点子。
不是做不了，是正经项目永远排在前面。
于是 commit 一拖再拖、代码里的刺没人看第二眼、写代码变成流水线。
我们想让这件事重新有点温度——所以，把仓库拟人化成一只仓鼠。"""),
    ("slide-03.png", """【机制 · 约 50 秒】
玩法很简单：输入 GitHub 用户名，只读公开仓库。
我们读三样东西：提交频率、commit message、最近一次 diff。
它们会变成四个状态：体重、心情、饥饿、亲密度。
一次 commit 等于一颗瓜子；7 天不写它 emo；14 天离家出走；
投喂一颗瓜子，它可能会回来。"""),
    ("slide-04.png", """【产品演示 · 约 40 秒】
产品打开就能玩。左边是领养入口，下面是它的窝、贡献图和状态。
门槛故意降到零：不用登录、不用授权。
现场可以把用户名给我，30 秒后你就能看到一只属于你的仓鼠，
以及它的第一句吐槽。右上角的 Kiro Hook 已连接——保存代码就能投喂。"""),
    ("slide-05.png", """【高光 · 约 40 秒】
它还有一个很好玩的腮帮子：读取最近一次 diff，给出三条轻量提示。
不是正式 code review，也不说教，用仓鼠的语气吐槽你。
比如 JSON.parse 没包 try/catch，它会说“崩了一下牙”。
它还会生成一张 4:5 的体检报告，保存、发群、发朋友圈都行。"""),
    ("slide-06.png", """【Kiro 深度 · 约 60 秒】
重点是：整个东西是 Kiro 做的。
Quick Spec 把一句话拆成需求、设计、任务；
并行 tasks 同时跑前端、后端和测试；
三个 Crew 子 Agent 负责仓鼠人格、commit 考古和界面；
最关键的一步是 Agent Hook：在 Kiro 里保存代码，
Hook 自动跑测试并 curl 一下 /feed，仓鼠当场吃瓜子。
Kiro Hook 不只是开发工具，它成了产品的消化系统。
多模型方面，Claude 写人格，DeepSeek 审代码，Auto 负责调度。"""),
    ("slide-07.png", """【为什么会被传播 · 约 30 秒】
为什么它会被转发？
第一，名字就是梗，仓库的仓、仓鼠的鼠，30 秒讲得清。
第二，每个人玩的是自己的仓库，看到的是自己的摆烂和卷。
第三，体检报告天生适合晒。
它没什么用，但它具备传播的所有要素。"""),
    ("slide-08.png", """【收尾 · 约 30 秒】
最后一句话：它没什么用，但你会因为它，想多写一行好代码。
这就是“无用之用”。
谢谢大家，欢迎现场领养，也欢迎问我 Kiro 到底怎么用。"""),
]


def add_fade(slide):
    sld = slide._element
    for t in sld.findall(qn("p:transition")):
        sld.remove(t)
    trans = etree.SubElement(sld, qn("p:transition"))
    trans.set("spd", "med")
    etree.SubElement(trans, qn("p:fade"))


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]
    for i, (img, note) in enumerate(SLIDES, 1):
        path = os.path.join(SHOTS, img)
        if not os.path.exists(path):
            raise SystemExit("missing " + path)
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        pic.name = "代码仓鼠·第%d页" % i
        # 演讲备注
        tf = slide.notes_slide.notes_text_frame
        tf.text = note
        # 淡入切换
        add_fade(slide)

    prs.save(OUT)

    # 自检
    check = Presentation(OUT)
    print("saved:", OUT)
    print("bytes:", os.path.getsize(OUT))
    print("slides:", len(check.slides))
    for i, s in enumerate(check.slides, 1):
        pics = [sh for sh in s.shapes if sh.shape_type == 13]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        trans = s._element.find(qn("p:transition"))
        print("  %d  images=%d  notes=%d字  transition=%s" % (
            i, len(pics), len(notes), "yes" if trans is not None else "no"))


if __name__ == "__main__":
    main()
